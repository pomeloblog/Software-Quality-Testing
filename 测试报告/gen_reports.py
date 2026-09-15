# -*- coding: utf-8 -*-
"""
按附录3（模块一测试报告）与附录4（模块二AI融合报告）模板生成全组测试报告。

数据来源：
  - 甲（王尚清）M1/M4：origin/wsq 分支 测试用例清单-M1-M4(甲).xlsx / 测试报告 / 缺陷报告
  - 丙 M3/M6：本仓库 测试用例/丙-* （12条用例实测：5 OK / 7 NG，8缺陷）
  - 乙 M2/M5：占位待补
用法：python gen_reports.py
"""
import copy
import os

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TPL_DIR = r"E:\软件质量测试\软件测试实践作业\实践作业-文档模板2026"
OUT1 = os.path.join(HERE, "测试报告（模块一）.docx")
OUT2 = os.path.join(HERE, "测试报告（模块二-AI融合实践）.docx")

# ============================================================
# 通用工具
# ============================================================

def is_toc_para(p):
    s = p.text
    return "\t" in s and s.strip()[-1:].isdigit()


def body_paras(doc):
    return [p for p in doc.paragraphs if not is_toc_para(p)]


def replace_body_text(doc, prefix, new_text, Occur=1):
    """把正文中以 prefix 开头的段落的文本替换为 new_text（保留首个 run 的格式）。"""
    n = 0
    for p in body_paras(doc):
        if p.text.strip().startswith(prefix):
            if not p.runs:
                p.add_run(new_text)
            else:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ""
            n += 1
            if n >= Occur:
                return p
    return None


def insert_table_after(anchor_para, doc, data, style="Table Grid", font_size=9):
    """在 anchor 段落后插入表格（doc.add_table 追加到文末后移动元素）。"""
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    try:
        t.style = style
    except Exception:
        pass
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(v))
            run.font.size = Pt(font_size)
    anchor_para._p.addnext(t._tbl)
    return t


def fill_cover(doc, system, extra_rev=()):
    for p in doc.paragraphs:
        s = p.text.strip()
        if s == "XX系统":
            if p.runs:
                p.runs[0].text = system
                for r in p.runs[1:]:
                    r.text = ""
    rev = doc.tables[0]
    rows = [("2026-09-15", "V1.0", "初稿：整合甲(M1/M4)、丙(M3/M6)实测数据；乙(M2/M5)待补充", "丙（统稿）")]
    rows += list(extra_rev)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            rev.cell(i + 1, j).text = str(v)


def update_toc(path):
    import win32com.client
    app = win32com.client.Dispatch("KWPS.Application")
    app.Visible = False
    d = app.Documents.Open(path)
    try:
        d.Fields.Update()
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).Update()
    except Exception as e:
        print("TOC update warn:", e)
    d.Save()
    d.Close(False)
    app.Quit()


# ============================================================
# 报告一：附录3 模块一测试报告
# ============================================================

OVERVIEW = (
    "被测系统为网上书店系统（Java/Servlet/JSP，MVC 三层架构：bean→dao→service→web，无构建工具，"
    "依赖 jar 位于 web/WEB-INF/lib）。系统分为前台（浏览/检索/购物车/下单）与后台（图书/用户/订单管理）"
    "两大部分，服务器 Tomcat 9（context path=/Book），数据库 MySQL（库名 book，4 张表：t_user、t_book、"
    "t_order、t_order_item），源码共 46 个 Java 文件。小组按功能内聚性选择 6 个关键模块开展测试："
)

MODULE_TBL = [
    ["模块", "核心功能", "主要源文件", "代码规模", "负责人"],
    ["M1 用户认证", "注册/登录/验证码/用户名查重", "UserServlet、UserServiceImpl、UserDaoImpl", "约400行", "甲（王尚清）"],
    ["M2 权限控制", "前后台隔离/越权访问", "ManagerFilter、TransactionFilter、BaseServlet", "约200行", "乙（待补）"],
    ["M3 图书浏览与检索", "分页/搜索/价格筛选/销量榜单", "ClientBookServlet、BookServiceImpl、BookDaoImpl", "约450行", "丙"],
    ["M4 购物车", "加购/累加/改量/删除/清空/金额", "CartServlet、Cart、CartItem", "约350行", "甲（王尚清）"],
    ["M5 订单与事务", "下单/库存扣减/事务提交回滚", "ClientOrderServlet、OrderServiceImpl", "约400行", "乙（待补）"],
    ["M6 后台管理", "图书/用户增删改查/发货/总账单", "BookServlet、ManagerUserServlet、ManagerOrderServlet", "约500行", "丙"],
]

SCOPE_TBL = [
    ["项目", "说明"],
    ["本次测试覆盖",
     "M1 用户认证（登录/注册/验证码/查重）；M3 图书浏览与检索（分页/搜索/价格筛选/榜单）；"
     "M4 购物车（加购/累加/改量/删除/清空/金额）；M6 后台管理（图书/用户增删改查/发货/总账单）。"
     "共 4 个模块 27 条已执行用例，全部自动化。"],
    ["本次测试不覆盖",
     "① M2 权限控制、M5 订单与事务——由小组成员乙负责，成果待其提交后汇总（用例计划 12 条）；"
     "② M1 的注销与个人信息修改、前端样式与交互视觉效果——优先级较低；"
     "③ 性能、并发、兼容性、安全性（除已发现的越权/验证码缺陷外）——超出本阶段（模块一：测试基础实践）范畴。"],
    ["不覆盖原因", "分工安排 / 优先级排序 / 阶段范围外 / 环境限制。"],
]

STRATEGY = (
    "黑盒方法按被测对象的输入域特征选择，全组共使用等价类划分、边界值分析、场景法三种（满足≥2种要求）：\n"
    "（1）等价类划分：用于输入域可明显划分合法/非法类的功能，如登录凭证（存在/不存在/密码对错）、"
    "用户名查重、新增图书（合法值/空值/负值/类型非法值）——从每类取代表值暴露校验缺失。\n"
    "（2）边界值分析：用于有数值边界或长度边界的参数，如分页 pageNo/pageSize（0、-1、超大、非数字）、"
    "价格区间 min/max（相等、交叉、小数）、购物车数量（0、负数、超大、非数字）——缺陷集中出现在边界处。\n"
    "（3）场景法：用于多步骤业务流程与权限路径，如管理员登录→后台管理全流程、搜索→翻页链接完整性、"
    "总账单对账、发货状态流转。\n"
    "白盒手段上，采用 JUnit 单元测试直调被测类（Cart/CartItem 金额计算、注册密码一致性），"
    "甲另以 JDK 动态代理 Stub 模拟 HttpServletRequest/Response 绕过验证码构造前置条件；"
    "其余用例均为 HTTP 集成测试打真实 Servlet，并以 JDBC 直连数据库断言落库结果（黑盒+数据库校验结合）。"
)

ENV_TEXT = (
    "测试在小组两台开发机上分别完成，被测系统与测试代码运行条件如下：\n"
    "机器A（甲）：Windows 11 专业版；Temurin JDK 1.8.0_504；Tomcat 9.0.121（context path=/Book）；"
    "MySQL 5.7.44（root/12345，见 src/jdbc.properties）；JUnit 4.12 + hamcrest 1.3。\n"
    "机器B（丙）：Windows 11 家庭版；JDK 11（测试进程）；Tomcat 9（context path=/Book，端口 8080）；"
    "MySQL 8.0（专用账号 bookstore/123456，因项目自带 mysql-connector 5.1.7 不支持 8.0 默认认证方式，"
    "创建 mysql_native_password 账号适配）；JUnit 4.12。\n"
    "被测应用账号：管理员 admin/admin；测试造数账号 cp_test_01 等（用后清理）。"
    "数据库由 sql/bookstore.sql 初始化（29 本图书、admin 用户、sales/stock 各 100）。"
    "两台机器版本不同（JDK8/11、MySQL5.7/8.0）但测试结果一致，说明被测系统对运行环境具备基本可移植性。"
)

CASE_OVERVIEW = [
    ["测试方法", "用例数量", "占比", "已自动化数"],
    ["等价类划分", "8", "29.6%", "8"],
    ["边界值分析", "6", "22.2%", "6"],
    ["场景法", "13", "48.1%", "13"],
    ["合计（已执行）", "27（乙M2/M5计划12条待补，总计39）", "—", "27（100%，≥24达标）"],
]

TYPICAL_CASES = [
    ("BS-IT-004 登录不校验验证码（甲，场景法）",
     "输入：正确用户名密码、不带验证码参数直接 POST /userServlet?action=login。"
     "预期：应校验 KAPTCHA_SESSION_KEY，验证码缺失时拒绝登录。"
     "实际：HTTP 200 直接登录成功——验证码形同虚设。设计依据：阅读 UserServlet.login() 发现只取 username/password，"
     "与 regist() 形成对照（后者校验验证码），属定向缺陷验证用例，可演示暴力破解风险。"),
    ("BS-IT-042 分页 pageSize 边界（丙，边界值法）",
     "输入：pageSize=6/abc/0/-1。预期：非法值容错为默认每页 4 条，不应服务器错误。"
     "实际：pageSize=0 返回 500（pageTotalCount%pageSize 除零 ArithmeticException）；-1 返回 500（LIMIT 负数 SQL 错误）。"
     "设计依据：BookServiceImpl.page() 中 pageSize 直接参与取模与 LIMIT，为典型边界缺陷点。"),
    ("BS-IT-044 搜索翻页链接未 URL 编码（丙，场景法）",
     "输入：nameorauthor=\"三 体\"（含空格）与 a&b（含连接符）搜索后检查翻页 href。"
     "预期：参数值 URL 编码，翻页后搜索条件完整。"
     "实际：href 中原始字符拼接（ClientBookServlet 第 48 行），点击下一页后搜索词截断、条件丢失。"
     "设计依据：代码走查发现分页 URL 拼接未编码，属缺陷定向验证。"),
    ("BS-IT-046 销量榜单跳过第一名（丙，缺陷定向）",
     "输入：造数将 id=1 图书销量改为全库最高（999），请求榜单后还原。"
     "预期：榜单含销量第一名。实际：不含——SQL 为 ORDER BY sales DESC LIMIT 1,50，偏移量 1 跳过第一名"
     "（BookDaoImpl 第 88 行）。设计依据：代码走查发现 LIMIT 常量可疑，通过造数使缺陷可见。"),
    ("BS-IT-102 新增图书输入校验（丙，等价类法）",
     "输入：合法图书 / 空书名 / 负价格 / 非数字价格四类代表值。"
     "预期：后三类被服务端校验拒绝。实际：全部直接入库（空名存空串、负价入库、abc 以 NULL 入库）。"
     "设计依据：add() 链路无任何校验，BeanUtils 转换异常被吞，属校验缺失型缺陷。"),
    ("BS-IT-065 购物车数量改负数（甲，边界值法）",
     "输入：updateCount&count=-5。预期：拒绝或修正为合法值。"
     "实际：数量 -5 生效，总数量与总金额为负，可进一步生成负金额订单。"
     "设计依据：Cart.updateCount 无下界校验，负数直接进入金额计算，危害数据正确性。"),
]

DIST_TBL = [
    ["序号", "被测模块", "需求简述", "测试用例数量"],
    ["1", "M1 用户认证", "登录凭证校验、验证码防绕过、注册合法性、用户名查重", "7"],
    ["2", "M2 权限控制", "前后台隔离、未登录/普通用户越权（乙负责，待汇总）", "5（待乙补充）"],
    ["3", "M3 图书浏览与检索", "分页参数容错、搜索模糊匹配、翻页条件保持、价格区间、榜单正确性", "6"],
    ["4", "M4 购物车", "加购累加、数量边界、金额精度、空车操作", "8"],
    ["5", "M5 订单与事务", "下单链路、库存边界、事务回滚（乙负责，待汇总）", "7（待乙补充）"],
    ["6", "M6 后台管理", "图书/用户增删改查合法性、回显正确性、发货流转、账单对账", "6"],
    ["合计", "—", "—", "27 已执行 / 39 计划"],
]

FRAMEWORK = (
    "测试框架：JUnit 4.12（项目自带 jar，不引入 Maven）。小组按成员分两个并行测试工程，入口均为一键脚本：\n"
    "工程A（甲，仓库根目录）：run-tests.bat + src/com/yj/test/{common, m1, m4}，"
    "含 HttpTestClient（自动管理 JSESSIONID）、MockHttp（JDK 动态代理桩，绕过验证码）、TestDb（JDBC 断言）。\n"
    "工程B（丙，tests/ 目录）：run-tests.bat + src/{HttpTestUtil, M3BookBrowseTest, M6ManagerTest} + lib/（自带 junit/hamcrest/mysql 驱动）。\n"
    "两层测试路线覆盖黑白盒：单元测试（白盒）直调被测方法（BS-UT-006/061/062）；"
    "集成测试（黑盒）经 HttpURLConnection 向运行中的 /Book 发真实 HTTP 请求并以 JDBC 断言数据库结果。"
)

KEY_IMPL = (
    "核心设计一：HTTP 会话保持与重定向处理（丙，HttpTestUtil，节选）：\n"
    "  // 手动跟随302：JDK自动跟随不传Cookie，登录页新发的JSESSIONID会污染会话（已踩坑）\n"
    "  for (int hop = 0; hop < 5; hop++) {\n"
    "      conn = (HttpURLConnection) new URL(url).openConnection();\n"
    "      conn.setInstanceFollowRedirects(false);\n"
    "      if (jsessionId != null) conn.setRequestProperty(\"Cookie\", \"JSESSIONID=\" + jsessionId);\n"
    "      int code = conn.getResponseCode();\n"
    "      String location = conn.getHeaderField(\"Location\");\n"
    "      if ((code == 302) && location != null) {\n"
    "          url = location.startsWith(\"http\") ? location\n"
    "              : location.startsWith(\"/\") ? SERVER_ROOT + location : BASE + \"/\" + location;\n"
    "          curMethod = \"GET\"; curBody = null; continue;   // 302后按浏览器惯例降级GET\n"
    "      }\n"
    "      return new Response(code, readAll(...));\n"
    "  }\n"
    "设计要点：(1)全部用例参数化集中在 URL 拼接层，用例方法只给输入与断言；"
    "(2)数据库断言与造数统一走 HttpTestUtil.exec/queryLong（预编译参数）；"
    "(3)改库用例以 try/finally + @AfterClass 双保险还原（如销量改999后还原、测试图书按作者标记清理），保证可重复执行；"
    "(4)外部依赖不 Mock 数据库——被测系统与断言共用同一真实库，验证真实落库；验证码等不可控依赖才用 Stub 替代。"
)

RESULT_TEXT = (
    "分模块执行结果如下（初测=缺陷发现阶段；甲的 M1/M4 已完成缺陷修复并回归）："
)

RESULT_TBL = [
    ["模块", "用例数", "初测通过", "初测失败", "初测通过率", "当前状态"],
    ["M1 用户认证", "7", "5", "2", "71.4%", "修复后回归 7/7 全通过"],
    ["M4 购物车", "8", "3", "5", "37.5%", "修复后回归 8/8 全通过"],
    ["M3 图书浏览与检索", "6", "2", "4", "33.3%", "8 个缺陷待修复（其中M3占5）"],
    ["M6 后台管理", "6", "3", "3", "50.0%", "8 个缺陷待修复（其中M6占3）"],
    ["M2/M5（乙）", "12", "待补充", "待补充", "—", "成果待乙提交"],
    ["已执行汇总", "27", "13", "14", "48.1%", "M1/M4 修复后整体 20/27=74.1%"],
]

DEFECT_TBL = [
    ["编号", "缺陷简述", "模块", "严重度", "发现方式", "状态", "发现人"],
    ["DEF-001", "登录接口不校验验证码，可绕过暴力破解", "M1", "高", "自动化", "已修复", "甲"],
    ["DEF-002", "购物车数量可改为负数，数量与金额为负", "M4", "高", "自动化", "已修复", "甲"],
    ["DEF-003", "非法数量参数被静默改为1，输入被篡改无提示", "M4", "高", "自动化", "已修复", "甲"],
    ["DEF-004", "注册不校验两次密码一致性", "M1", "中", "自动化", "已修复", "甲"],
    ["DEF-005", "加购不存在图书ID空指针，HTTP 500", "M4", "中", "自动化", "已修复", "甲"],
    ["DEF-006", "加购非数字ID同样 HTTP 500", "M4", "中", "自动化", "已修复", "甲"],
    ["DEF-007", "空购物车删除/改量时无任何响应", "M4", "中", "自动化", "已修复", "甲"],
    ["BUG-M3-01", "pageSize=0/-1 服务端500（除零/SQL错误）", "M3", "高", "自动化", "打开", "丙"],
    ["BUG-M3-02", "搜索翻页链接未URL编码，条件丢失", "M3", "高", "自动化", "打开", "丙"],
    ["BUG-M3-03", "价格区间 min>max 无校验无提示", "M3", "中", "自动化", "打开", "丙"],
    ["BUG-M3-04", "价格参数仅支持整数，小数静默失效", "M3", "中", "自动化", "打开", "丙"],
    ["BUG-M3-05", "销量榜单 LIMIT 1,50 跳过第一名", "M3", "中", "自动化", "打开", "丙"],
    ["BUG-M6-01", "新增图书无校验：空名/负价/NULL价入库", "M6", "高", "自动化", "打开", "丙"],
    ["BUG-M6-02", "删除图书非数字id 500；不存在id静默", "M6", "中", "自动化", "打开", "丙"],
    ["BUG-M6-03", "新增重复用户名唯一约束冲突 500", "M6", "高", "自动化", "打开", "丙"],
]

ROOT_CAUSE = (
    "闭环一（DEF-001，登录不校验验证码，高/安全）："
    "①发现：自动化用例 BS-IT-004 不带验证码直接 POST 登录，返回 200 且登录成功；"
    "②定位：对照 UserServlet.login()（仅取 username/password）与 regist()（校验 KAPTCHA_SESSION_KEY）；"
    "③根因：登录流程从未读取 session 中的验证码，验证码仅是前端装饰；"
    "④修复：login() 增加 Kaptcha 校验（与注册一致），为兼容自动化引入可配置开关；"
    "⑤验证：回归 BS-IT-004 不带验证码被拒，带正确验证码可登录，其余 6 条 M1 用例无回归。\n"
    "闭环二（DEF-002，购物车数量可为负，高/数据正确性）："
    "①发现：BS-IT-065 以 count=-5 调 updateCount，返回 302 且 totalCount 变为 -3；"
    "②定位：Cart.updateCount() 直接 setCount(count)，无下界校验；"
    "③根因：数量直接进入金额计算，负数导致总数量/总金额为负，可下负金额订单；"
    "④修复：updateCount 对 count<1 时忽略或修正为 1 并提示；"
    "⑤验证：回归 BS-IT-065 负数不再生效，066 非数字提示明确，金额计算用例（062）无回归。\n"
    "闭环三（BUG-M3-05，榜单跳过第一名，中/功能正确性——丙，待修复闭环）："
    "①发现：BS-IT-046 造数后榜单不含销量第一名；②定位：queryForPageItemsOrder() SQL 为 LIMIT 1,50；"
    "③根因：LIMIT 偏移量 1 跳过第一行；④修复建议：改为 LIMIT 50（一行改动）；"
    "⑤验证方案：修复后重跑 run-tests.bat，test046 应通过且榜首为销量最高图书。"
)

QUALITY = (
    "M1 用户认证与 M4 购物车：7 个缺陷全部修复并通过 15 条用例回归（通过率 100%），正常路径实现质量良好，达到发布标准。\n"
    "M3 图书浏览与检索、M6 后台管理：12 条用例通过 5 条（41.7%），8 个缺陷（高4/中4）处于打开状态，"
    "缺陷集中于参数校验缺失与异常处理缺失，其中翻页 URL 未编码、新增图书无校验影响核心用户体验，未达发布标准；"
    "建议优先修复 4 个高级别缺陷（其中榜单修复仅需一行），修复后以 run-tests.bat 回归。\n"
    "M2 权限控制、M5 订单与事务：成果待乙提交，整体结论待补。"
    "综合判断：被测系统在已测 4 模块的正常功能路径质量可接受，但异常输入防御普遍薄弱（15 个缺陷中 11 个属校验缺失），"
    "整体暂不建议发布，待高级别缺陷清零后复评。"
)

CONTRIB = [
    ["成员", "承担模块", "具体任务", "测试代码行数", "发现缺陷数", "工作量占比"],
    ["甲（王尚清）", "M1 用户认证\nM4 购物车", "15条用例设计与自动化；HTTP客户端/动态代理桩/TestDb基础设施；7缺陷修复与回归；M1/M4用例清单、缺陷报告、测试报告初稿", "约1194行", "7", "33.3%"],
    ["乙（待填）", "M2 权限控制\nM5 订单与事务", "12条用例设计与自动化；缺陷报告（待提交后更新本表）", "（待填）", "（待填）", "33.3%"],
    ["丙", "M3 图书浏览与检索\nM6 后台管理", "测试环境搭建（MySQL8适配/Tomcat部署/Git仓库）；12条用例设计与自动化（run-tests.bat一键运行）；8缺陷定位至代码行；全组报告统稿", "约700行", "8", "33.3%"],
    ["合计", "6 个模块", "", "", "15", "100%"],
]


def build_report1():
    doc = Document(os.path.join(TPL_DIR, "附录3：测试报告模板（模块一）.docx"))
    fill_cover(doc, "网上书店系统")
    # 先捕获模板表格引用（后续插入新表会使 doc.tables 索引偏移）
    scope_t, overview_t, dist_t = doc.tables[1], doc.tables[2], doc.tables[3]

    # 1 概述：替换占位段 + 模块表
    p = replace_body_text(doc, "说明被测软件系统名称", OVERVIEW)
    insert_table_after(p, doc, MODULE_TBL)

    # 2.1 范围表
    for i, row in enumerate(SCOPE_TBL):
        if i == 0:
            continue
        for j, v in enumerate(row):
            scope_t.cell(i, j).text = v

    # 2.2 策略
    replace_body_text(doc, "在此说明针对不同测试项", STRATEGY)

    # 2.3 环境
    p = replace_body_text(doc, "在此说明程序的运行条件", ENV_TEXT)
    # 删除第二个占位段（"并说明测试代码的运行条件"）已并入上文
    replace_body_text(doc, "并说明测试代码的运行条件", "")

    # 3.1 用例总览表（需加行）
    t = overview_t
    for i, row in enumerate(CASE_OVERVIEW):
        if i == 0:
            continue
        if i == 1:
            for j, v in enumerate(row):
                t.cell(1, j).text = v
        else:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = v

    # 3.2 典型用例
    txt = "\n".join(f"{i+1}. {t_}" for i, t_ in enumerate(TYPICAL_CASES))
    replace_body_text(doc, "针对不同被测模块，共选取", txt)

    # 3.3 用例分布表（需加行）
    t = dist_t
    for i, row in enumerate(DIST_TBL):
        if i == 0:
            continue
        if i <= 2:
            for j, v in enumerate(row):
                t.cell(i, j).text = v
        else:
            cells = t.add_row().cells
            for j, v in enumerate(row):
                cells[j].text = v

    # 4.1 / 4.2
    replace_body_text(doc, "测试框架：如 pytest", FRAMEWORK)
    replace_body_text(doc, "选取核心的测试代码片段", KEY_IMPL)

    # 5.1 执行结果（文字+表格）
    p = replace_body_text(doc, "分不同的功能模块，说明总用例数", RESULT_TEXT)
    insert_table_after(p, doc, RESULT_TBL)

    # 5.2 缺陷摘要
    p = replace_body_text(doc, "针对每个缺陷，简要描述", f"初测共确认有效缺陷 15 个（高 7 / 中 8），其中 7 个已修复（甲），8 个待修复（丙）。摘要如下：")
    insert_table_after(p, doc, DEFECT_TBL)

    # 5.3 根因分析
    replace_body_text(doc, "选取2-3个最有代表性的缺陷", ROOT_CAUSE)

    # 5.4 质量结论
    replace_body_text(doc, "基于测试结果，得出被测模块", QUALITY)

    # 6 贡献度：模板只有占位段，插入表格
    p = replace_body_text(doc, "说明每个成员承担的具体任务", "各成员工作量如下（乙的数据待其提交后更新）：")
    insert_table_after(p, doc, CONTRIB)

    doc.save(OUT1)
    print("saved:", OUT1)


# ============================================================
# 报告二：附录4 模块二 AI 融合报告（方案2：测试手段含AI）
# ============================================================

AI_OVERVIEW = (
    "被测系统为网上书店系统（JavaWeb MVC，46 个 Java 文件，Tomcat 9 + MySQL），小组选择 6 个模块中的 4 个"
    "（M1/M3/M4/M6，另 M2/M5 由乙人工完成）开展测试。本模块二实践采用方案2（测试手段含 AI）："
    "在传统测试流程中系统性引入 AI 编程助手辅助完成环境搭建、测试用例设计、自动化脚本开发、缺陷定位与文档编写，"
    "并对 AI 产出进行人工复核与修正。选择方案 2 的理由：被测对象为传统确定性 JavaWeb 系统（无 AI 模块），"
    "而测试过程中存在大量重复性机械劳动（环境适配、脚本编写、文档整理），是 AI 辅助增效的典型场景；"
    "同时小组成员可完整观察\"AI 产出—人工复核—修正\"的协作过程，为反思提供真实素材。"
    "使用工具：Claude Code（命令行 AI 编程助手，Anthropic），全程对话记录留存于会话日志。"
)

AI_PROCESS = (
    "实践按\"AI 执行→人工复核→修正定稿\"循环推进，关键过程如下：\n"
    "（1）环境搭建：AI 完成 Git 仓库初始化与 GitHub 连接、MySQL 8.0 认证适配（识别出项目自带 mysql-connector 5.1.7 "
    "与 8.0 默认认证不兼容，创建 mysql_native_password 专用账号）、编译部署到 Tomcat 并验证登录流程。"
    "人工复核：确认数据库导入结果与页面数据一致。【此处插入截图1：环境搭建对话，待补】\n"
    "（2）用例设计：AI 依据分工文档的模块定义与代码走查结果，生成 M3/M6 共 12 条用例（边界值 4、等价类 3、场景法 5），"
    "编号 BS-IT-041~046/101~106，按模板字段输出 Excel。人工修正：用例取舍与预期结果的产品语义判断（如\"删除不存在的 id "
    "是否必须提示\"由人拍板）。【此处插入截图2：用例生成对话，待补】\n"
    "（3）脚本开发：AI 实现 HttpTestUtil/M3/M6 测试类与 run-tests.bat 一键运行；实测 12 条得 5 OK / 7 NG。\n"
    "（4）缺陷定位：AI 将 7 条 NG 用例全部定位到具体代码行（如 BookServiceImpl 第 59 行除零、"
    "BookDaoImpl 第 88 行 LIMIT 1,50），并按附录2 模板生成 8 个缺陷的缺陷报告。"
    "人工复核：对关键缺陷用 curl 独立复现确认。【此处插入截图3：缺陷定位对话，待补】\n"
    "（5）文档统稿：AI 按附录2/3/4 模板生成缺陷报告与测试报告并更新目录域。"
)

AI_RESULTS = (
    "新增测试成果：M3/M6 共 12 条自动化用例（100% 自动化）与 8 个实测确认缺陷（高 4 / 中 4），"
    "全部有可复现步骤与代码级根因，详见 测试用例/丙-测试用例清单-M3-M6.xlsx 与 丙-缺陷报告-M3-M6.docx。\n"
    "AI 辅助与人工改进的关键差异对比：\n"
    "①效率：环境搭建+脚本开发+三份文档，AI 辅助下约半天完成，估计纯人工需 2-3 天；\n"
    "②深度：AI 主动走查源码发现的缺陷（榜单 LIMIT 1、翻页未编码）超出用例清单原始线索；\n"
    "③准确性：AI 首版测试工具存在会话保持缺陷导致 1 条用例假通过（详见第 3 节），经人工发现异常数据后复查修正——"
    "修正后该用例（BS-IT-102）才暴露出真实缺陷，说明 AI 产出的测试代码本身也必须被验证；\n"
    "④判断力：严重程度定级、用例取舍、模板符合性等需要人对课程要求的理解，AI 仅能给出建议。"
)

AI_REFLECT = (
    "AI 工具帮我们做了什么：环境适配（MySQL 8 认证兼容方案）、46 个源文件的快速走查、12 条用例设计与三份模板文档生成、"
    "8 个缺陷的根因定位（精确到代码行）与修复建议、一键测试脚本（run-tests.bat）。\n"
    "AI 工具哪里\"不靠谱\"（真实案例）：①HttpTestUtil 首版使用 JDK 自动跟随 302 重定向，不传 Cookie，"
    "被登录页新发的 JSESSIONID 污染会话，导致 BS-IT-102（新增图书无校验）连续两轮假通过——AI 自己的测试代码引入了新的缺陷模式，"
    "直到人工发现\"用例通过但数据库无新增记录\"的矛盾才推动排查；②生成缺陷报告时按标题前缀查找正文段落，"
    "误匹配到目录条目，将\"1 引言\"标题覆盖为结果分析内容（已修正为过滤目录段）；③Excel 首版列数与模板不符（14 列 vs 13 列）。"
    "共性：AI 错误多发生在\"对既有约定（模板结构、HTTP 语义）的隐性假设\"上，且错误输出表面通常\"看起来很对\"。\n"
    "人做了什么：发现用例结果与数据库状态的矛盾并要求复查（假通过的唯一线索来自人的交叉核对）；"
    "用 curl 独立复现关键缺陷；决定用例预期结果的产品语义（提示 vs 静默、定级高/中）；核对模板符合性与提交规范。\n"
    "AI 工具在测试中的角色定位：辅助者。判断依据：AI 在机械性、模式化劳动上（编译部署、脚本、文档、代码走查）"
    "效率与覆盖面显著优于人工，但其产出必须经过人工验证才能采信——本次实践中 AI 的两处错误（会话污染、目录误匹配）"
    "均由人工复核发现，且\"预期结果是否符合产品语义\"的判断始终需要人承担最终责任。"
)

AI_CONTRIB = [
    ["成员", "承担任务", "工作量占比"],
    ["甲（王尚清）", "M1/M4 用例与自动化（人工为主）、7 缺陷修复回归、模块一报告 M1/M4 部分", "33.3%"],
    ["乙（待填）", "M2/M5 用例与自动化（待提交后更新）", "33.3%"],
    ["丙", "AI 辅助实践的组织与复核、M3/M6 全部 AI 协作产出（12 用例/8缺陷/三份文档）、模块二报告撰写", "33.3%"],
    ["合计", "", "100%"],
]


def build_report2():
    doc = Document(os.path.join(TPL_DIR, "附录4：测试报告模板（模块二）.docx"))
    fill_cover(doc, "网上书店系统")

    p = replace_body_text(doc, "说明被测软件系统名称", AI_OVERVIEW)
    replace_body_text(doc, "并说明选择方案1", "选择方案2（测试手段含AI），理由见上段。")
    replace_body_text(doc, "方案1：描述被测AI模块", "")
    replace_body_text(doc, "方案2：描述使用的AI工具", "方案2：使用工具为 Claude Code（命令行 AI 编程助手），应用场景覆盖测试环境搭建、测试用例生成、自动化测试脚本开发、缺陷定位与测试文档编写。")

    replace_body_text(doc, "方案1：AI模块的测试用例设计", AI_PROCESS)
    replace_body_text(doc, "方案2：AI的原始输出", "")
    replace_body_text(doc, "方案1：AI相关测试用例清单", AI_RESULTS)

    replace_body_text(doc, "在此说明程序的运行条件", "与模块一测试报告 2.3 节一致（两台开发机：JDK8/Tomcat9.0.121/MySQL5.7.44 与 JDK11/Tomcat9/MySQL8.0；AI 工具运行于本机命令行，无额外服务依赖）。")
    replace_body_text(doc, "并说明测试代码的运行条件", "")

    replace_body_text(doc, "方案1：反思AI系统本身的质量", AI_REFLECT)
    # 列表占位段（方案1的4条）置空
    replace_body_text(doc, "被测AI做对了什么", "")
    replace_body_text(doc, "被测AI做错了什么", "")
    replace_body_text(doc, "测试AI系统的挑战", "")
    replace_body_text(doc, "AI系统能否通过测试证明", "")

    p = replace_body_text(doc, "说明每个成员承担的具体任务", "模块二实践中各成员工作量如下（乙待更新）：")
    insert_table_after(p, doc, AI_CONTRIB)

    doc.save(OUT2)
    print("saved:", OUT2)


if __name__ == "__main__":
    build_report1()
    build_report2()
    update_toc(OUT1)
    update_toc(OUT2)
    print("done")
