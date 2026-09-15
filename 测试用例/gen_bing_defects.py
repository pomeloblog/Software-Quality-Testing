# -*- coding: utf-8 -*-
"""
按附录2《缺陷报告模板.doc》生成丙的缺陷报告（M3+M6，8个缺陷）。

流程：
  1. WPS COM 把模板 .doc 另存为 .docx（保留原版式）
  2. python-docx 填充：封面/修订表/引言/环境/测试章节 + 逐缺陷表格（深拷贝模板示例表）
  3. WPS COM 更新目录域后保存

依赖：本机 WPS/Word（COM）、pywin32、python-docx。用法：python gen_bing_defects.py
"""
import copy
import os

TPL = r"E:\软件质量测试\软件测试实践作业\实践作业-文档模板2026\附录2：缺陷报告模板.doc"
HERE = os.path.dirname(os.path.abspath(__file__))
TMP_DOCX = os.path.join(HERE, "_template_converted.docx")
OUT = os.path.join(HERE, "丙-缺陷报告-M3-M6.docx")

M3 = "M3 图书浏览与检索"
M6 = "M6 后台管理"

# 每个缺陷的字段（对应模板18行缺陷表）
DEFECTS_M3 = [
    dict(no="001", case="BS-IT-042", item="图书分页", demand="分页参数应容错非法取值",
         title="分页参数pageSize为0或负数时服务端500错误", sev="高", prio="高",
         desc=["1. GET /Book/client/bookServlet?action=page&pageNo=1&pageSize=0 返回HTTP 500；pageSize=-1同样500。",
               "2. 复现：浏览器或curl直接访问上述URL即可，无需登录。",
               "3. 根因：pageSize=0时BookServiceImpl.page()中pageTotalCount%pageSize抛除零ArithmeticException；"
               "pageSize=-1时pageTotal为负→pageNo归一为负→LIMIT起始/条数为负→MySQL语法错误（BookServiceImpl.java第59、63-68行）。"],
         expect="非法pageSize应容错为默认每页4条，正常返回列表页（HTTP 200）。",
         actual="HTTP 500服务器错误页。",
         related="无",
         fix="已修复：BookServiceImpl的page()/pageByPrice()/pageByNameOrAuthor()三个方法入口统一增加 if (pageSize < 1) pageSize = Page.PAGE_SIZE; 容错。",
         note="修复验证（已执行）：run-tests.bat复测test042通过——pageSize=0/-1/abc均返回HTTP 200且每页4条。"),
    dict(no="002", case="BS-IT-044", item="图书搜索", demand="搜索翻页应保持搜索条件完整传递",
         title="搜索结果翻页链接未URL编码，含空格/&的搜索条件丢失", sev="高", prio="高",
         desc=["1. 访问 /Book/client/bookServlet?action=pageByNameOrAuthor&nameorauthor=%E4%B8%89%20%E4%BD%93（搜索\"三 体\"）。",
               "2. 页面翻页链接href形如 client/bookServlet?action=pageByNameOrAuthor&nameorauthor=三 体&pageNo=2，空格未编码。",
               "3. 点击\"下一页\"后搜索词被截断为\"三\"；含&的搜索词（如a&b）后段被解析为独立参数，搜索条件丢失。",
               "4. 根因：ClientBookServlet.java第46-49行将原始参数值直接拼接进分页URL，未调用URLEncoder.encode。"],
         expect="翻页链接对参数值做URL编码（空格→%20、&→%26），点击后搜索条件完整传递。",
         actual="翻页后搜索条件丢失，结果与无条件列表一致或错误。",
         related="无",
         fix="已修复：ClientBookServlet翻页URL拼接处改为 URLEncoder.encode(nameorauthor, \"UTF-8\") 后再 append。",
         note="修复验证（已执行）：test044通过——翻页链接形如 nameorauthor=%E4%B8%89+%E4%BD%93，空格与&不再截断参数。"),
    dict(no="003", case="BS-IT-045", item="价格筛选", demand="价格区间参数应做合法性校验",
         title="价格区间min>max无校验无提示，静默返回空列表", sev="中", prio="中",
         desc=["1. 访问 /Book/client/bookServlet?action=pageByPrice&min=100&max=10。",
               "2. 页面显示空列表，无任何错误或提示，用户无从得知区间设置错误。",
               "3. 根因：BETWEEN 100 AND 10恒为假，ClientBookServlet.pageByPrice()与BookServiceImpl对min>max未做任何校验。"],
         expect="给出可见提示（如\"价格区间设置有误\"）或自动交换min/max归一化处理。",
         actual="静默空列表。",
         related="BUG-M3-004（同一用例BS-IT-045的另一处参数缺陷）",
         fix="已修复：pageByPrice对min>max自动交换两值，并在index.jsp以红色提示显示\"价格区间设置有误（最小值大于最大值），已自动调整区间\"。",
         note="修复验证（已执行）：test045通过——min=100&max=10返回提示且展示调整后区间图书。"),
    dict(no="004", case="BS-IT-045", item="价格筛选", demand="价格参数应支持小数（价格为DECIMAL类型）",
         title="价格筛选参数仅支持整数，小数价格静默失效", sev="中", prio="中",
         desc=["1. 访问 /Book/client/bookServlet?action=pageByPrice&min=56.5&max=56.5（库中《三体》恰为56.5元）。",
               "2. 返回的不是该书，而是全部图书的第一页——筛选条件被完全忽略。",
               "3. 根因：WebUtils.parseInt按int解析\"56.5\"抛NumberFormatException，被catch后静默返回默认值min=0、max=Integer.MAX_VALUE，等价于取消筛选（WebUtils.java parseInt）。"],
         expect="小数价格参数正常参与筛选，精确命中56.5元的书。",
         actual="筛选失效，返回全部图书。",
         related="BUG-M3-003（同一用例BS-IT-045）",
         fix="已修复：WebUtils新增parseDouble()；ClientBookServlet价格参数按double解析；BookService/BookDao/BookDaoImpl的pageByPrice链路参数类型由int改为double。",
         note="修复验证（已执行）：test045通过——min=56.5&max=56.5精确命中《三体》（1条）。"),
    dict(no="005", case="BS-IT-046", item="销量榜单", demand="榜单应包含销量第一名且降序排列",
         title="销量榜单SQL使用LIMIT 1,50，跳过销量第一名", sev="中", prio="中",
         desc=["1. 造数：UPDATE t_book SET sales=999 WHERE id=1（《解忧杂货店》销量改为全库最高）。",
               "2. 访问 /Book/client/bookServlet?action=pageOrder，榜单中找不到《解忧杂货店》，从原第二名开始展示。",
               "3. 测试后还原：UPDATE t_book SET sales=100 WHERE id=1。",
               "4. 根因：BookDaoImpl.java第88行 queryForPageItemsOrder()的SQL为 ORDER BY sales DESC LIMIT 1,50，偏移量1跳过第一行。"],
         expect="榜单包含销量第一名，按销量降序排列。",
         actual="第一名缺失。",
         related="无",
         fix="已修复：BookDaoImpl.queryForPageItemsOrder()的SQL由 LIMIT 1,50 改为 LIMIT 50。",
         note="修复验证（已执行）：test046通过——造数后销量第一名《解忧杂货店》出现在榜单首位。"),
]

DEFECTS_M6 = [
    dict(no="006", case="BS-IT-102", item="图书管理-新增", demand="新增图书应校验必填与数值合法性",
         title="新增图书无服务端输入校验：空书名/负价格/非法价格直接入库", sev="高", prio="高",
         desc=["1. admin登录后台，POST /Book/manager/bookServlet?action=add 新增图书。",
               "2. 书名留空提交→新增成功，库中该行name=''；价格提交-5→库中price=-5.00；价格提交abc→BeanUtils转换异常被WebUtils吞掉后以NULL入库。",
               "3. 全流程无任何输入校验。根因：BookServlet.add()无校验；WebUtils.copyParamToBean()仅printStackTrace吞异常。"],
         expect="空书名、负价格、非数字价格应被服务端校验拒绝并给出提示。",
         actual="三类非法输入全部直接写库。",
         related="无",
         fix="已修复：BookServlet新增validateBook()（书名非空、价格≥0且合法数字、销量/库存≥0），add/update入库前调用，不合法转发回book_edit.jsp显示红色提示；非法价格不再以NULL入库。",
         note="修复验证（已执行）：test102通过——空书名/负价格/非数字价格均被拒绝且库中无脏数据，合法图书正常入库。"),
    dict(no="007", case="BS-IT-104", item="图书管理-删除", demand="删除接口应容错非法与不存在的id",
         title="删除图书：非数字id引发500；不存在的id静默无提示", sev="中", prio="中",
         desc=["1. admin登录后访问 /Book/manager/bookServlet?action=delete&id=abc → HTTP 500（Integer.parseInt抛NumberFormatException未处理）。",
               "2. 访问 action=delete&id=999999（不存在）→ 正常跳回列表页，无任何提示，用户不知道删除未生效。",
               "3. 根因：BookServlet.delete()第34-35行parseInt未防护、删除影响行数未检查。"],
         expect="非数字id提示参数错误；不存在的id提示\"图书不存在/删除失败\"。",
         actual="非数字id为500错误页；不存在id静默成功假象。",
         related="无",
         fix="已修复：delete改用WebUtils.parseInt容错（非法id取-1不再抛异常）；deleteBookById改为返回影响行数，0行时重定向携带\"删除失败：图书不存在\"，book_manager.jsp顶部显示提示。",
         note="修复验证（已执行）：test104通过——id=abc与id=999999均返回200并显示\"图书不存在\"提示，存在的id正常删除。"),
    dict(no="008", case="BS-IT-105", item="用户管理-新增", demand="新增用户应处理重复用户名",
         title="新增重复用户名触发唯一约束，服务端500错误页", sev="高", prio="高",
         desc=["1. admin登录后POST /Book/manager/UserServlet?action=add，参数username=admin&password=x&email=dup@test.com。",
               "2. 返回Tomcat 500错误页；日志出现SQLIntegrityConstraintViolationException: Duplicate entry 'admin' for key 't_user.username'。",
               "3. 根因：ManagerUserServlet.add()调用链无重复检查、无异常处理，异常经TransactionFilter包装后抛给容器。"],
         expect="提示\"用户名已存在\"并停留在新增页（与前台注册的AJAX查重体验一致）。",
         actual="HTTP 500错误页，事务回滚。",
         related="无",
         fix="已修复：ManagerUserServlet.add入库前调用userService.existsUsername预检（并补用户名非空校验），重复时转发回user_edit.jsp显示\"新增失败：用户名已存在\"，不再触发唯一约束异常。",
         note="修复验证（已执行）：test105通过——重复用户名返回200带提示，正常新增/删除不受影响。"),
]

# ---------------- 第1步：WPS COM 转换 .doc → .docx ----------------

def convert_template():
    import win32com.client
    app = win32com.client.Dispatch("KWPS.Application")
    app.Visible = False
    doc = app.Documents.Open(TPL, ReadOnly=True)
    doc.SaveAs2(TMP_DOCX, FileFormat=12)  # wdFormatXMLDocument
    doc.Close(False)
    app.Quit()


# ---------------- 第2步：python-docx 填充 ----------------

def set_cell(cell, lines):
    """单元格写入多行文本（列表）。"""
    if isinstance(lines, str):
        lines = [lines]
    cell.text = ""
    first = True
    for ln in lines:
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        run = p.add_run(ln)
        run.font.size = cell.paragraphs[0].runs[0].font.size if p.runs else run.font.size


def is_toc_para(p):
    """目录条目段落：文本以制表符+页码结尾（如 '3.4.1  管理员\t6'）。"""
    s = p.text
    return "\t" in s and s.strip()[-1:].isdigit()


def body_paras(doc):
    return [p for p in doc.paragraphs if not is_toc_para(p)]


def find_para(doc, prefix):
    for p in body_paras(doc):
        if p.text.strip().startswith(prefix):
            return p
    return None


def fill_defect_table(table, d, module, tester="丙"):
    t = table
    set_cell(t.cell(0, 1), tester)
    set_cell(t.cell(0, 3), "2026-09-15")
    set_cell(t.cell(1, 1), module)
    set_cell(t.cell(1, 3), module.split(" ")[0])
    set_cell(t.cell(2, 1), d["item"])
    set_cell(t.cell(3, 1), d["demand"])
    set_cell(t.cell(4, 1), d["case"])
    set_cell(t.cell(5, 1), d["sev"])
    set_cell(t.cell(5, 3), d["prio"])
    set_cell(t.cell(5, 5), "关闭")
    set_cell(t.cell(6, 1), "丙（开发+测试同一人，课程作业）")
    set_cell(t.cell(7, 1), "项目组全体成员")
    set_cell(t.cell(8, 1), d["title"])
    set_cell(t.cell(9, 0), d["desc"] + ["", "预期结果：", d["expect"], "", "实际结果：", d["actual"]])
    set_cell(t.cell(10, 1), "无（运行tests\\run-tests.bat输出可复现）")
    set_cell(t.cell(11, 1), d["related"])
    set_cell(t.cell(12, 1), d["note"])
    # 解决区（已实施修复并回归）
    set_cell(t.cell(14, 1), "丙")                # 解决者
    set_cell(t.cell(14, 3), "2026-09-15")        # 解决日期
    set_cell(t.cell(15, 1), "v1.1-fix")          # 解决build
    set_cell(t.cell(15, 3), "已修复并回归通过")   # 解决方案
    set_cell(t.cell(16, 1), d["fix"])
    set_cell(t.cell(17, 1), "丙（回归 run-tests.bat 12/12 通过后关闭）")  # 关闭者


def main():
    convert_template()

    from docx import Document
    doc = Document(TMP_DOCX)

    # --- 封面 ---
    for p in doc.paragraphs:
        s = p.text.strip()
        if s == "待测软件名称":
            p.runs[0].text = "待测软件名称：网上书店系统（JavaWeb MVC）"
            for r in p.runs[1:]:
                r.text = ""
        elif "XX测试缺陷报告书" in p.text:
            for r in p.runs:
                r.text = r.text.replace("XX测试缺陷报告书", "功能测试缺陷报告书")

    # --- 修订表 ---
    rev = doc.tables[0]
    rows = [("2026-09-14", "1.0", "初稿：M3+M6共12条自动化用例与首轮实测", "丙"),
            ("2026-09-15", "1.1", "修复测试工具会话保持问题后复测，确认8个缺陷", "丙"),
            ("2026-09-15", "1.2", "8个缺陷全部修复（v1.1-fix）并回归12/12通过，缺陷全部关闭", "丙")]
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            set_cell(rev.cell(i + 1, j), v)

    # --- 1.2 背景 ---
    fills = {
        "本软件名称：": "本软件名称：网上书店系统（Java/Servlet/JSP，MVC架构，Tomcat部署，MySQL存储）",
        "本项目的任务提出者：": "本项目的任务提出者：《软件测试与质量保证实践》课程（武剑洁）",
        "开发者：": "开发者：网上书店原项目作者（开源学习项目）；测试小组第三组",
        "用户：": "用户：书店顾客（浏览、购物车、下单）与书店管理员（图书/用户/订单管理）",
    }
    for p in doc.paragraphs:
        s = p.text.strip()
        if s in fills:
            p.runs[0].text = fills[s]
            for r in p.runs[1:]:
                r.text = ""

    # --- 1.3 定义 ---
    p = find_para(doc, "（此处主要对本文档中提到的一些术语进行解释")
    if p:
        p.runs[0].text = ("OK/POK/NG/NT：用例执行结果代号（通过/部分通过/不通过/用例有误）。"
                          "BS-IT-nnn：集成测试用例编号规则（本项目）。"
                          "M3/M6：图书浏览与检索模块/后台管理模块。")

    # --- 1.4 参考资料表 ---
    ref = doc.tables[1]
    refs = [("1", "网上书店系统项目说明", "网上书店项目/readme.md", "2022-09", "开源项目"),
            ("2", "测试模块划分与分工", "测试模块划分与分工.md", "2026-09", "课程小组"),
            ("3", "实践作业文档模板2026", "软件测试实践作业/实践作业-文档模板2026", "2026-09", "课程")]
    for i, row in enumerate(refs):
        for j, v in enumerate(row):
            set_cell(ref.cell(i + 1, j), v)

    # --- 2 测试环境 ---
    p = find_para(doc, "CPU、硬盘、内存")
    if p:
        p.runs[0].text = "普通PC即可（本次测试：CPU 12代酷睿、内存16GB、SSD），无特殊硬件要求。"
    p = find_para(doc, "操作系统，开发平台")
    if p:
        p.runs[0].text = ("操作系统 Windows 11；JDK 11（测试进程）；Tomcat 9.0（context path=/Book，端口8080）；"
                          "MySQL 8.0（库名book，账号bookstore）；JUnit 4.12 + hamcrest 1.3；"
                          "测试脚本仓库tests/目录，run-tests.bat一键执行。")

    # --- 3 测试章节标题与正文 ---
    p = find_para(doc, "3冒烟测试")
    if p:
        p.runs[0].text = "3  功能测试（自动化集成测试）"
        for r in p.runs[1:]:
            r.text = ""
    p = find_para(doc, "（此处应说明被测软件的版本")
    if p:
        p.runs[0].text = ("被测对象为网上书店系统，源码位于仓库 网上书店项目/，"
                          "编译后部署于本机Tomcat 9（http://localhost:8080/Book/），"
                          "数据库由 sql/bookstore.sql 初始化（29本图书、admin账号）。")
    p = find_para(doc, "在各类用户的功能中")
    if p:
        p.runs[0].text = ("本次测试针对M3图书浏览与检索、M6后台管理两个模块的全部12条用例展开"
                          "（BS-IT-041~046、BS-IT-101~106），优先级为高的测试项全覆盖；"
                          "方法采用边界值法、等价类法与场景法结合，全部用例经JUnit自动化执行。")
    p = find_para(doc, "（此处应具体说明执行的过程")
    if p:
        p.runs[0].text = ("1. 启动MySQL并导入bookstore.sql；2. 编译项目并部署到Tomcat（/Book）；"
                          "3. 运行 tests\\run-tests.bat 一键编译并执行12条自动化用例；"
                          "4. 对NG用例用curl/浏览器人工复核，确认缺陷后登记于3.4节。")

    # --- 3.4 缺陷表 ---
    # 章节小标题 3.4.1 管理员 → M3
    p341 = find_para(doc, "3.4.1")
    if p341:
        p341.runs[0].text = "3.4.1  M3 图书浏览与检索"
        for r in p341.runs[1:]:
            r.text = ""
    p = find_para(doc, "缺陷如表所示")
    if p:
        p.runs[0].text = "M3模块确认缺陷5个，如表3.4.1~3.4.5所示。"

    # 模板示例表 = tables[2]；先深拷贝一份空白模板
    blank_tbl = copy.deepcopy(doc.tables[2]._tbl)
    example_caption = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("表3.4.1"):
            example_caption = copy.deepcopy(p._p)
            break

    all_groups = [("M3 图书浏览与检索", DEFECTS_M3), ("M6 后台管理", DEFECTS_M6)]

    # 先把 3.4.2 标题插入到最后一个M3表格之后，再逐表生成
    from docx.oxml.ns import qn

    # 填第一张表（模板自带，对应M3缺陷001）
    caption_texts = []
    fill_defect_table(doc.tables[2], DEFECTS_M3[0], M3)

    anchor = doc.tables[2]._tbl
    table_no = 1

    def insert_after(anchor_el, new_el):
        anchor_el.addnext(new_el)
        return new_el

    # 为 M3 其余4个 + M6 3个 插入 [标题段落 + 表格]
    rest = [(M3, d) for d in DEFECTS_M3[1:]] + [(M6, d) for d in DEFECTS_M6]
    first_m6 = True
    for module, d in rest:
        if first_m6 and module == M6:
            # 插入 3.4.2 标题（拷贝3.4.1标题段落样式）
            h = copy.deepcopy(p341._p)
            anchor = insert_after(anchor, h)
            from docx.text.paragraph import Paragraph
            Paragraph(h, doc).runs[0].text = "3.4.2  M6 后台管理"
            first_m6 = False
        # 题注段落
        cap = copy.deepcopy(example_caption)
        anchor = insert_after(anchor, cap)
        table_no += 1
        d["cap_text"] = f"表3.4.{table_no}  {d['no']}-{d['title']}"
        from docx.text.paragraph import Paragraph
        pr = Paragraph(cap, doc)
        pr.runs[0].text = d["cap_text"]
        for r in pr.runs[1:]:
            r.text = ""
        # 空白表
        tbl_el = copy.deepcopy(blank_tbl)
        anchor = insert_after(anchor, tbl_el)
        # 找到刚插入的表并填充
        from docx.table import Table
        fill_defect_table(Table(tbl_el, doc), d, module)

    # 第一张表题注改写（正文中的题注段落；目录条目会被COM重新生成，无需处理）
    from docx.text.paragraph import Paragraph
    for p in body_paras(doc):
        if p.text.strip().startswith("表3.4.1"):
            p.runs[0].text = f"表3.4.1  {DEFECTS_M3[0]['no']}-{DEFECTS_M3[0]['title']}"
            for r in p.runs[1:]:
                r.text = ""
            break

    # --- 3.5 结果分析和结论 ---
    p = None
    found35 = False
    for para in body_paras(doc):
        if para.text.strip().startswith("3.5"):
            found35 = True
        elif found35 and para.text.strip():
            p = para
            break
    if p:
        p.runs[0].text = (
            "本次功能测试共执行12条自动化用例（M3模块6条、M6模块6条），初测通过5条、不通过7条，"
            "确认有效缺陷8个：严重程度高4个（pageSize除零500、翻页URL未编码、新增图书无校验、重复用户名500），"
            "中4个（min>max无提示、小数价格失效、榜单跳过第一名、删除id容错缺失）。"
            "缺陷分布规律明显：M3集中于参数校验缺失与SQL细节错误，M6集中于输入校验与异常处理缺失，"
            "反映出项目缺少统一的参数校验层和全局异常处理。\n"
            "修复闭环：8个缺陷已于2026-09-15全部修复（build v1.1-fix），"
            "修复后运行 tests\\run-tests.bat 回归，12条用例全部通过（通过率100%），无回归缺陷，"
            "数据库测试数据全部还原。被测的M3、M6两个模块达到发布标准，全部缺陷关闭。")

    doc.save(OUT)

    # --- 第3步：更新目录域 ---
    import win32com.client
    app = win32com.client.Dispatch("KWPS.Application")
    app.Visible = False
    d = app.Documents.Open(OUT)
    try:
        d.Fields.Update()
        for i in range(1, d.TablesOfContents.Count + 1):
            d.TablesOfContents(i).Update()
    except Exception as e:
        print("TOC update warn:", e)
    d.Save()
    d.Close(False)
    app.Quit()

    os.remove(TMP_DOCX)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
